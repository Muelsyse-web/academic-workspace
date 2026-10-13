"""Bounded raster assets from immutable PPTX originals; no layout reconstruction."""
import hashlib
from io import BytesIO
from pathlib import Path
import re
from uuid import uuid4
import warnings

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import FileResponse
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .parsing import _validate_office_archive

router = APIRouter(prefix='/api/v1')
MAX_IMAGE_PIXELS = 20_000_000
MAX_IMAGE_BYTES = 16 * 1024 * 1024
MAX_SLIDE_IMAGES = 50
MAX_SOURCE_BYTES = 100 * 1024 * 1024
ALLOWED_FORMATS = {'PNG', 'JPEG', 'WEBP'}


def _safe_path(root: Path, *parts: str) -> Path:
    """Reject symlinks at every storage component, including a cache root."""
    path = root
    for part in parts:
        path = path / part
        if path.is_symlink():
            raise HTTPException(422, 'Preview storage is unavailable')
    if not path.resolve().is_relative_to(root.resolve()):
        raise HTTPException(422, 'Preview storage is unavailable')
    return path


def _source(db, identifier):
    doc = db.one('SELECT kind,blob_hash FROM documents WHERE id=?', (identifier,))
    if not doc:
        raise HTTPException(404, 'Document not found')
    if doc['kind'] != 'pptx':
        raise HTTPException(422, 'Slide images are available for PPTX documents only')
    checksum = doc['blob_hash']
    if not re.fullmatch(r'[0-9a-f]{64}', checksum):
        raise HTTPException(422, 'Invalid source checksum')
    source = _safe_path(db.root, 'originals', checksum)
    if not source.is_file():
        raise HTTPException(404, 'Original missing; restore from backup')
    if source.stat().st_size > MAX_SOURCE_BYTES:
        raise HTTPException(422, 'Original exceeds the 100 MiB preview limit')
    with source.open('rb') as stream:
        actual = hashlib.file_digest(stream, 'sha256').hexdigest()
    if actual != checksum:
        raise HTTPException(422, 'Original checksum mismatch; restore from backup')
    return source, checksum


def _pictures(shapes):
    # Iteration avoids Python recursion limits for nested picture groups.
    pending = [iter(shapes)]
    pictures = []
    while pending:
        shape = next(pending[-1], None)
        if shape is None:
            pending.pop()
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            pending.append(iter(shape.shapes))
        elif shape._element.tag.endswith('}pic'):
            pictures.append(shape)
            if len(pictures) > MAX_SLIDE_IMAGES:
                raise HTTPException(422, 'Slide exceeds the preview limit of 50 image assets')
    return pictures


def _dimensions(image):
    width, height = image.size
    if width <= 0 or height <= 0 or width * height > MAX_IMAGE_PIXELS:
        raise HTTPException(422, 'Slide image exceeds the 20 million pixel preview limit')
    return width, height


def _cache_image(shape, target):
    if target.is_file():
        if target.stat().st_size > MAX_IMAGE_BYTES:
            raise HTTPException(422, 'Cached preview exceeds the 16 MiB image limit')
        with Image.open(target, formats=['PNG']) as image:
            return _dimensions(image)
    blob = shape.image.blob
    if len(blob) > MAX_IMAGE_BYTES:
        raise HTTPException(422, 'Slide image exceeds the 16 MiB preview limit')
    with warnings.catch_warnings():
        warnings.simplefilter('error', Image.DecompressionBombWarning)
        with Image.open(BytesIO(blob)) as image:
            if image.format not in ALLOWED_FORMATS:
                raise HTTPException(422, 'Unsupported slide image; previews accept PNG, JPEG, or WebP only')
            dimensions = _dimensions(image)
            image.load()
            # Re-encode pixels only: strip embedded metadata and never serve source bytes.
            pixels = image.convert('RGBA' if 'A' in image.getbands() or 'transparency' in image.info else 'RGB')
            output = BytesIO()
            pixels.save(output, format='PNG')
    if output.tell() > MAX_IMAGE_BYTES:
        raise HTTPException(422, 'Decoded preview exceeds the 16 MiB image limit')
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_name(f'.{uuid4().hex}.tmp')
    try:
        with temporary.open('xb') as stream:
            stream.write(output.getvalue())
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)
    return dimensions


def _assets(db, identifier, slide, requested_index=None):
    if slide < 1 or (requested_index is not None and requested_index < 0):
        raise HTTPException(404, 'Slide or image not found')
    try:
        # Coordinate with restore/maintenance and concurrent cache writers.
        with db.maintenance:
            source, checksum = _source(db, identifier)
            _validate_office_archive(source)
            deck = Presentation(source)
            if slide > len(deck.slides):
                raise HTTPException(404, 'Slide not found')
            pictures = _pictures(deck.slides[slide - 1].shapes)
            if requested_index is not None and requested_index >= len(pictures):
                raise HTTPException(404, 'Slide image not found')
            assets = []
            indices = range(len(pictures)) if requested_index is None else [requested_index]
            for index in indices:
                target = _safe_path(db.root, 'media-cache', checksum, str(slide), f'{index}.png')
                width, height = _cache_image(pictures[index], target)
                assets.append(({'index': index, 'width': width, 'height': height,
                                'url': f'/api/v1/documents/{identifier}/slides/{slide}/images/{index}',
                                'alt': f'Slide {slide} image {index + 1}'}, target))
            return assets
    except HTTPException:
        raise
    except Exception as exc:
        # Library diagnostics may contain local paths or source content.
        raise HTTPException(422, 'Unable to preview this slide: invalid or unsupported PPTX image data') from exc


@router.get('/documents/{identifier}/slides/{slide}/images')
def slide_images(identifier: str, slide: int, request: Request):
    return [metadata for metadata, _ in _assets(request.app.state.db, identifier, slide)]


@router.get('/documents/{identifier}/slides/{slide}/images/{index}')
def slide_image(identifier: str, slide: int, index: int, request: Request):
    _, target = _assets(request.app.state.db, identifier, slide, index)[0]
    return FileResponse(target, media_type='image/png',
                        headers={'X-Content-Type-Options': 'nosniff', 'Cache-Control': 'private, no-cache'})
