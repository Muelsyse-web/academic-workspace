"""Extract immutable source positions without inferring Office page numbers."""
from pathlib import Path, PurePosixPath
import logging
import re
import stat
import zipfile

TEXT_BLOCK_LINES = 40
TEXT_BLOCK_CHARS = 8000
MAX_ARCHIVE_ENTRIES = 10000
MAX_ARCHIVE_BYTES = 256 * 1024 * 1024


def _validate_office_archive(path):
    """Budget ZIP metadata before an Office library expands any member."""
    with path.open("rb") as source:
        if source.read(8) == bytes.fromhex("d0cf11e0a1b11ae1"):
            raise ValueError("Encrypted or legacy Office containers are unsupported; supply an unlocked DOCX/PPTX.")
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise ValueError("Office archive exceeds 10000 entries.")
            expanded = 0
            for entry in entries:
                name = entry.filename
                parts = PurePosixPath(name.replace("\\", "/"))
                if (parts.is_absolute() or ".." in parts.parts or "\\" in name
                        or re.match(r"^[A-Za-z]:", name)
                        or stat.S_ISLNK(entry.external_attr >> 16)):
                    raise ValueError("Unsafe path or traversal in Office archive.")
                if entry.flag_bits & 1:
                    raise ValueError("Encrypted Office archives are unsupported.")
                expanded += entry.file_size
                if expanded > MAX_ARCHIVE_BYTES:
                    raise ValueError("Office archive expanded size exceeds the allowed budget.")
                if entry.file_size > 1024 * 1024 and entry.file_size / max(entry.compress_size, 1) > 1000:
                    raise ValueError("Office archive compression ratio exceeds 1000.")
    except zipfile.BadZipFile as exc:
        raise ValueError("Invalid or corrupt Office archive.") from exc


def _unit(locator_type, locator, heading, text):
    return {"locator_type": locator_type, "locator": locator,
            "heading": heading[:300], "text": text.strip()}


def _first_line(text, fallback):
    return next((line.strip() for line in text.splitlines() if line.strip()), fallback)


def _pdf_units(path):
    from pypdf import PdfReader

    # pypdf diagnostics can include source bytes. Keep this dependency's
    # namespace out of application logs; callers report sanitized parse errors.
    pdf_logger = logging.getLogger("pypdf")
    pdf_logger.setLevel(logging.CRITICAL + 1)
    if not any(isinstance(handler, logging.NullHandler) for handler in pdf_logger.handlers):
        pdf_logger.addHandler(logging.NullHandler())
    pdf_logger.propagate = False

    units = []
    with path.open("rb") as source:
        reader = PdfReader(source)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF documents are unsupported; supply an unlocked copy.")
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            units.append(_unit("page", number, _first_line(text, path.stem), text))
    return units


def _docx_units(path):
    from docx import Document
    from docx.table import Table

    document = Document(path)
    heading = path.stem
    heading_path = []
    units = []
    for number, block in enumerate(document.iter_inner_content(), 1):
        if isinstance(block, Table):
            text = "\n".join("\t".join(cell.text for cell in row.cells) for row in block.rows)
        else:
            text = block.text
            style = block.style.name.lower() if block.style else ""
            if text.strip() and (style.startswith("heading") or style == "title"):
                match = re.search(r"(\d+)$", style)
                level = int(match.group(1)) if match else (0 if style == "title" else 1)
                heading_path = [(depth, title) for depth, title in heading_path if depth < level]
                heading_path.append((level, text.strip()))
                heading = " / ".join(title for _, title in heading_path)
        units.append(_unit("paragraph", number, heading, text))
    return units


def _slide_text(shapes):
    parts = []
    for shape in shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if shape.has_table:
            parts.append("\n".join("\t".join(cell.text for cell in row.cells) for row in shape.table.rows))
        if hasattr(shape, "shapes"):
            parts.extend(_slide_text(shape.shapes))
    return parts


def _pptx_units(path):
    from pptx import Presentation

    deck = Presentation(path)
    units = []
    for number, slide in enumerate(deck.slides, 1):
        text = "\n".join(_slide_text(slide.shapes))
        title = slide.shapes.title
        heading = title.text.strip() if title is not None else ""
        units.append(_unit("slide", number, heading or _first_line(text, path.stem), text))
    return units


def _text_units(path, markdown):
    try:
        content = path.read_text(encoding="utf-8-sig")
    except UnicodeError as exc:
        raise ValueError("Text must use UTF-8 encoding.") from exc
    lines = content.splitlines()
    heading = path.stem
    units, block = [], []
    start, block_chars = 1, 0

    def flush():
        nonlocal block, block_chars
        if block:
            units.append(_unit("line", start, heading, "\n".join(block)))
            block, block_chars = [], 0

    for number, line in enumerate(lines, 1):
        match = re.match(r"^#{1,6}\s+(.+?)\s*#*\s*$", line) if markdown else None
        if match:
            flush()
            heading = match.group(1)
        elif number == 1 and line.strip():
            heading = line.strip()
        if len(line) > TEXT_BLOCK_CHARS:
            flush()
            for offset in range(0, len(line), TEXT_BLOCK_CHARS):
                units.append(_unit("line", number, heading, line[offset:offset + TEXT_BLOCK_CHARS]))
            continue
        if block and (len(block) >= TEXT_BLOCK_LINES or block_chars + 1 + len(line) > TEXT_BLOCK_CHARS):
            flush()
        if not block:
            start = number
        block_chars += len(line) + (1 if block else 0)
        block.append(line)
    flush()
    return units or [_unit("line", 1, heading, "")]


def parse_document(path: Path, kind: str) -> list[dict]:
    """Return text units with positive, physical source positions."""
    path = Path(path)
    if kind not in {"pdf", "pptx", "docx", "txt", "md"}:
        raise ValueError(f"Unsupported document type: {kind}")
    try:
        if kind in {"txt", "md"}:
            return _text_units(path, kind == "md")
        if kind in {"docx", "pptx"}:
            _validate_office_archive(path)
        if kind == "docx":
            return _docx_units(path)
        if kind == "pdf":
            return _pdf_units(path)
        return _pptx_units(path)
    except OSError as exc:
        raise ValueError("Unable to read document.") from exc
    except (ValueError, ImportError):
        raise
    except Exception as exc:
        # Third-party readers use different exception families for malformed
        # PDF streams, XML, relationships, and package metadata.
        raise ValueError(f"Invalid or corrupt {kind.upper()} document.") from exc
