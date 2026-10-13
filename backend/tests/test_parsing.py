"""Synthetic files only: source positions, extraction, and input boundaries."""
import importlib
import importlib.util
import logging
from pathlib import Path
import subprocess
import sys
import zipfile

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from reportlab.pdfgen import canvas


def parse(path, kind):
    if importlib.util.find_spec("workspace.parsing") is None:
        pytest.fail("parse_document is not implemented yet")
    return importlib.import_module("workspace.parsing").parse_document(path, kind)


def test_pdf_preserves_real_pages_and_empty_pages(tmp_path):
    path = tmp_path / "lecture.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 720, "Probability foundations")
    pdf.drawString(72, 690, "Bayes rule connects conditional probabilities.")
    pdf.showPage()
    pdf.showPage()
    pdf.drawString(72, 720, "Linear algebra")
    pdf.save()
    units = parse(path, "pdf")
    assert [u["locator"] for u in units] == [1, 2, 3]
    assert {u["locator_type"] for u in units} == {"page"}
    assert units[0]["heading"] == "Probability foundations"
    assert "Bayes rule" in units[0]["text"]
    assert units[1]["text"] == ""
    assert units[2]["heading"] == "Linear algebra"


def test_pptx_preserves_slides_titles_tables_and_empty_slide(tmp_path):
    path = tmp_path / "slides.pptx"
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[5])
    first.shapes.title.text = "机器学习 Machine learning"
    table = first.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Meaning"
    table.cell(1, 0).text = "Gradient 梯度"
    table.cell(1, 1).text = "Direction of change"
    deck.slides.add_slide(deck.slide_layouts[6])
    third = deck.slides.add_slide(deck.slide_layouts[5])
    third.shapes.title.text = "Eigenvectors"
    deck.save(path)
    units = parse(path, "pptx")
    assert [u["locator"] for u in units] == [1, 2, 3]
    assert {u["locator_type"] for u in units} == {"slide"}
    assert units[0]["heading"] == "机器学习 Machine learning"
    assert "Gradient 梯度\tDirection of change" in units[0]["text"]
    assert units[1]["text"] == ""
    assert units[2]["text"] == "Eigenvectors"


def test_docx_counts_paragraphs_and_tables_in_document_order(tmp_path):
    path = tmp_path / "notes.docx"
    doc = Document()
    doc.add_heading("Probability 概率", level=1)
    doc.add_paragraph("Conditional probability 条件概率")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "P(A|B)"
    table.cell(0, 1).text = "P(A and B) / P(B)"
    doc.add_paragraph("")
    doc.add_heading("Linear algebra", level=2)
    doc.add_paragraph("Av = lambda v")
    doc.save(path)
    units = parse(path, "docx")
    assert [u["locator"] for u in units] == [1, 2, 3, 4, 5, 6]
    assert {u["locator_type"] for u in units} == {"paragraph"}
    assert units[2]["text"] == "P(A|B)\tP(A and B) / P(B)"
    assert units[2]["heading"] == "Probability 概率"
    assert units[5]["heading"] == "Probability 概率 / Linear algebra"


@pytest.mark.parametrize("kind", ["txt", "md"])
def test_text_blocks_have_bounded_size_and_actual_start_lines(tmp_path, kind):
    path = tmp_path / ("notes." + kind)
    lines = ["# 概率 Probability"] + [f"Line {n}: synthetic lesson" for n in range(2, 86)]
    path.write_text("\n".join(lines), encoding="utf-8")
    units = parse(path, kind)
    assert [u["locator"] for u in units] == [1, 41, 81]
    assert {u["locator_type"] for u in units} == {"line"}
    assert "Line 41:" in units[1]["text"]
    assert all(len(u["text"].splitlines()) <= 40 for u in units)
    assert "概率 Probability" in units[0]["heading"]


def test_long_text_lines_are_bounded_without_losing_content(tmp_path):
    path = tmp_path / "long.txt"
    path.write_text("x" * 20000 + "\nsecond line", encoding="utf-8")
    units = parse(path, "txt")
    assert all(len(u["text"]) <= 8000 for u in units)
    assert sum(u["text"].count("x") for u in units) == 20000
    assert [u["locator"] for u in units[:3]] == [1, 1, 1]
    assert units[-1]["locator"] == 2


@pytest.mark.parametrize("kind", ["pdf", "pptx", "docx"])
def test_corrupt_document_has_helpful_value_error(tmp_path, kind):
    path = tmp_path / ("broken." + kind)
    path.write_bytes(b"not a legal document")
    with pytest.raises(ValueError, match="(?i)(invalid|corrupt|read)"):
        parse(path, kind)


def test_pdf_diagnostics_do_not_log_document_bytes(tmp_path, caplog):
    path = tmp_path / "private-header.pdf"
    path.write_bytes(b"SYNTHETIC_PRIVATE_HEADER_NOT_A_PDF")
    with caplog.at_level(logging.WARNING):
        with pytest.raises(ValueError, match="Invalid or corrupt PDF") as failure:
            parse(path, "pdf")
        logging.getLogger("workspace.parser-test").warning("application-diagnostic-still-visible")
    assert "SYNTH" not in caplog.text
    assert "SYNTH" not in str(failure.value)
    assert "application-diagnostic-still-visible" in caplog.text


def test_encrypted_pdf_is_rejected(tmp_path):
    path = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    writer.write(path)
    with pytest.raises(ValueError, match="(?i)encrypted"):
        parse(path, "pdf")


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_encrypted_or_legacy_office_container_has_helpful_error(tmp_path, kind):
    path = tmp_path / ("encrypted." + kind)
    path.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + b"\x00" * 504)
    with pytest.raises(ValueError, match="(?i)(encrypted|legacy)"):
        parse(path, kind)


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_office_zip_missing_document_parts_has_value_error(tmp_path, kind):
    path = tmp_path / ("incomplete." + kind)
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "broken XML")
    with pytest.raises(ValueError, match="(?i)(invalid|corrupt)"):
        parse(path, kind)


@pytest.mark.parametrize("member", ["../outside.xml", "/absolute.xml", "word/../../outside.xml", "C:\\outside.xml", "word\\..\\outside.xml"])
def test_office_archive_rejects_traversal_before_parsing(tmp_path, member):
    path = tmp_path / "unsafe.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(member, "malicious entry")
    with pytest.raises(ValueError, match="(?i)(unsafe|traversal)"):
        parse(path, "docx")


def test_office_archive_entry_count_limit(tmp_path):
    path = tmp_path / "too-many.pptx"
    with zipfile.ZipFile(path, "w") as archive:
        for n in range(10001):
            archive.writestr(f"entry{n}.xml", "")
    with pytest.raises(ValueError, match="(?i)(10000|entries)"):
        parse(path, "pptx")


def test_office_archive_compression_ratio_limit(tmp_path):
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (2 * 1024 * 1024))
    with pytest.raises(ValueError, match="(?i)(ratio|compression)"):
        parse(path, "docx")


def test_office_archive_expanded_size_limit(tmp_path, monkeypatch):
    if importlib.util.find_spec("workspace.parsing") is None:
        pytest.fail("archive size validation is not implemented yet")
    parsing = importlib.import_module("workspace.parsing")
    monkeypatch.setattr(parsing, "MAX_ARCHIVE_BYTES", 100)
    path = tmp_path / "large.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("one.xml", "a" * 60)
        archive.writestr("two.xml", "b" * 60)
    with pytest.raises(ValueError, match="(?i)(expanded|size)"):
        parse(path, "docx")


def test_unsupported_and_non_utf8_text_are_rejected(tmp_path):
    path = tmp_path / "binary.txt"
    path.write_bytes(b"\xff\xfe\x00")
    with pytest.raises(ValueError, match="(?i)unsupported"):
        parse(path, "exe")
    with pytest.raises(ValueError, match="(?i)(UTF-8|encoding)"):
        parse(path, "txt")


def test_fixture_generator_creates_five_legal_formats_inside_project(tmp_path):
    project = Path(__file__).resolve().parents[2]
    script = project / "scripts/generate_fixtures.py"
    assert script.exists(), "Synthetic fixture generator is not implemented yet"
    result = subprocess.run([sys.executable, str(script), "--output", str(tmp_path / "demo")], cwd=project, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    files = list((tmp_path / "demo").iterdir())
    assert {p.suffix for p in files} == {".pdf", ".pptx", ".docx", ".txt", ".md"}
    for path in files:
        units = parse(path, path.suffix[1:])
        assert len(units) >= 2
        assert any("Probability" in u["text"] for u in units)


def test_fixture_generator_rejects_output_outside_project(tmp_path):
    project = Path(__file__).resolve().parents[2]
    script = project / "scripts/generate_fixtures.py"
    assert script.exists(), "Synthetic fixture generator is not implemented yet"
    result = subprocess.run([sys.executable, str(script), "--output", str(project.parent / "outside-fixtures")], cwd=project, capture_output=True, text=True)
    assert result.returncode != 0
    assert "within" in result.stderr.lower()
