"""Raster slide previews preserve originals and respect untrusted-file budgets."""
import hashlib
from io import BytesIO
from pathlib import Path
import zipfile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import pytest

from workspace.db import Database, now


def deck_bytes(image_format='PNG', size=(48, 32), count=1, grouped=False):
    image = BytesIO()
    Image.new('RGB', size, '#236a86').save(image, format=image_format)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = 'Synthetic slide text'
    slide.shapes.add_table(1, 1, 0, Inches(1), Inches(2), Inches(1)).table.cell(0, 0).text = 'Table text'
    shapes = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
    for _ in range(count):
        shapes.add_picture(BytesIO(image.getvalue()), 0, 0, width=Inches(1))
    deck.slides.add_slide(deck.slide_layouts[6])
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


@pytest.fixture
def preview_app(tmp_path):
    from workspace.previews import router
    db = Database(tmp_path / 'data')
    db.initialize()
    app = FastAPI()
    app.state.db = db
    app.include_router(router)
    with TestClient(app) as client:
        yield client, db


def store(db, content, kind='pptx', blob_hash=None):
    checksum = blob_hash or hashlib.sha256(content).hexdigest()
    root = db.root / 'originals'
    root.mkdir(exist_ok=True)
    if '/' not in checksum:
        (root / checksum).write_bytes(content)
    with db.connect() as con:
        con.execute('INSERT INTO blobs VALUES (?,?,?,?)', (checksum, len(content), kind, now()))
        con.execute('INSERT INTO documents(id,blob_hash,title,filename,kind,created_at) VALUES (?,?,?,?,?,?)',
                    ('example', checksum, 'Synthetic', 'synthetic.' + kind, kind, now()))
    return root / checksum


def url(slide=1):
    return f'/api/v1/documents/example/slides/{slide}/images'


@pytest.mark.parametrize('image_format', ['PNG', 'JPEG'])
def test_real_images_have_dimensions_safe_bytes_and_stable_cache(preview_app, image_format):
    client, db = preview_app
    source = store(db, deck_bytes(image_format=image_format, grouped=True))
    original = source.read_bytes()
    response = client.get(url())
    assert response.status_code == 200
    assert response.json() == [{'index': 0, 'width': 48, 'height': 32,
                                'url': url() + '/0', 'alt': 'Slide 1 image 1'}]
    image_response = client.get(response.json()[0]['url'])
    assert image_response.status_code == 200
    assert image_response.headers['content-type'] == 'image/png'
    assert Image.open(BytesIO(image_response.content)).size == (48, 32)
    cache = db.root / 'media-cache' / source.name / '1' / '0.png'
    assert cache.is_file()
    first_modified = cache.stat().st_mtime_ns
    assert client.get(url()).json() == response.json()
    assert cache.stat().st_mtime_ns == first_modified
    assert source.read_bytes() == original
    assert hashlib.sha256(original).hexdigest() == source.name
    assert str(db.root) not in response.text
    assert client.get(url(2)).json() == []


@pytest.mark.parametrize('path', [url(0), url(-1), url(3), url() + '/-1', url() + '/1',
                                '/api/v1/documents/missing/slides/1/images'])
def test_unknown_slide_image_or_document_returns_not_found(preview_app, path):
    client, db = preview_app
    store(db, deck_bytes())
    assert client.get(path).status_code == 404


@pytest.mark.parametrize('kind,content', [('txt', b'plain text'), ('pptx', b'not an archive')])
def test_wrong_kind_or_corrupt_office_rejected(preview_app, kind, content):
    client, db = preview_app
    store(db, content, kind=kind)
    response = client.get(url())
    assert response.status_code == 422
    assert str(db.root) not in response.text


@pytest.mark.parametrize('mutation', ['checksum', 'missing', 'unsafe_hash', 'symlink'])
def test_invalid_source_rejected_even_after_cache_exists(preview_app, tmp_path, mutation):
    client, db = preview_app
    if mutation == 'unsafe_hash':
        store(db, deck_bytes(), blob_hash='../escape')
    else:
        source = store(db, deck_bytes())
        assert client.get(url()).status_code == 200
        if mutation == 'checksum':
            source.write_bytes(b'changed')
        elif mutation == 'missing':
            source.unlink()
        else:
            other = tmp_path / 'other.pptx'
            source.replace(other)
            source.symlink_to(other)
    response = client.get(url())
    assert response.status_code in (404, 422)
    assert str(db.root) not in response.text


@pytest.mark.parametrize('limit', ['pixels', 'bytes', 'count', 'format'])
def test_image_budget_and_format_rejection(preview_app, monkeypatch, limit):
    from workspace import previews
    client, db = preview_app
    if limit == 'pixels':
        monkeypatch.setattr(previews, 'MAX_IMAGE_PIXELS', 100)
    if limit == 'bytes':
        monkeypatch.setattr(previews, 'MAX_IMAGE_BYTES', 10)
    store(db, deck_bytes(count=51 if limit == 'count' else 1,
                         image_format='BMP' if limit == 'format' else 'PNG'))
    response = client.get(url())
    assert response.status_code == 422
    assert str(db.root) not in response.text
    assert not list((db.root / 'media-cache').rglob('*.png'))
    if limit == 'count':
        assert '50' in response.json()['detail']


def test_archive_safety_checked_before_presentation_decode(preview_app):
    client, db = preview_app
    output = BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        archive.writestr('../private-path', 'unsafe')
    store(db, output.getvalue())
    response = client.get(url())
    assert response.status_code == 422
    assert 'private-path' not in response.text


def test_cache_symlink_cannot_escape_data_root(preview_app, tmp_path):
    client, db = preview_app
    source = store(db, deck_bytes())
    external = tmp_path / 'external'
    external.mkdir()
    (db.root / 'media-cache').symlink_to(external, target_is_directory=True)
    response = client.get(url())
    assert response.status_code == 422
    assert list(external.iterdir()) == []
    assert source.is_file()


def replace_embedded_image(content, replacement):
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(content)) as original, zipfile.ZipFile(output, 'w') as edited:
        for entry in original.infolist():
            edited.writestr(entry, replacement if entry.filename.startswith('ppt/media/') else original.read(entry))
    return output.getvalue()


def test_webp_embedded_pixels_are_normalized_to_png(preview_app):
    client, db = preview_app
    raster = BytesIO()
    Image.new('RGB', (13, 17), 'red').save(raster, format='WEBP')
    store(db, replace_embedded_image(deck_bytes(), raster.getvalue()))
    response = client.get(url())
    assert response.status_code == 200
    assert (response.json()[0]['width'], response.json()[0]['height']) == (13, 17)
    assert client.get(url() + '/0').content.startswith(b'\x89PNG')


def test_corrupt_embedded_image_has_sanitized_error(preview_app):
    client, db = preview_app
    store(db, replace_embedded_image(deck_bytes(), b'private embedded content'))
    response = client.get(url())
    assert response.status_code == 422
    assert 'private embedded content' not in response.text


def test_synthetic_fixture_includes_real_png(tmp_path):
    import runpy
    fixture_script = Path(__file__).resolve().parents[2] / 'scripts' / 'generate_fixtures.py'
    runpy.run_path(str(fixture_script))['generate'](tmp_path / 'fixtures')
    deck = Presentation(tmp_path / 'fixtures' / 'synthetic-foundations.pptx')
    pictures = [shape for slide in deck.slides for shape in slide.shapes if shape._element.tag.endswith('}pic')]
    assert pictures, 'Synthetic PPTX needs a real embedded image to exercise previews'
    assert Image.open(BytesIO(pictures[0].image.blob)).format == 'PNG'
