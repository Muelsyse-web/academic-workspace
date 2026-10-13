#!/usr/bin/env python3
"""Generate only fictional, bilingual academic documents inside this project."""
import argparse
from io import BytesIO
from pathlib import Path


LESSONS = [
    ("Machine learning 机器学习", [
        "Gradient descent 梯度下降 minimizes a differentiable objective.",
        "Update: theta_next = theta - learning_rate * gradient.",
        "Synthetic example: f(x) = x^2, gradient at x = 3 is 6.",
    ]),
    ("Probability 概率", [
        "Conditional probability 条件概率: P(A|B) = P(A and B) / P(B).",
        "Bayes rule 贝叶斯公式: P(A|B) = P(B|A) * P(A) / P(B).",
        "Synthetic example: P(A and B) = 0.2 and P(B) = 0.5, so P(A|B) = 0.4.",
    ]),
    ("Linear algebra 线性代数", [
        "An eigenvector 特征向量 satisfies Av = lambda * v.",
        "For A = diag(2, 3), vector (1, 0) has eigenvalue 2.",
        "An orthogonal matrix 正交矩阵 satisfies Q^T Q = I.",
    ]),
]


def generate(output: Path):
    # Imports are local so path validation does not depend on document libraries.
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from PIL import Image, ImageDraw
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    output.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdf = canvas.Canvas(str(output / "synthetic-foundations.pdf"))
    pdf.setTitle("Synthetic academic foundations")
    for heading, lines in LESSONS:
        pdf.setFont("STSong-Light", 17)
        pdf.drawString(40, 790, heading)
        pdf.setFont("STSong-Light", 11)
        for index, line in enumerate(lines):
            pdf.drawString(40, 750 - index * 35, line)
        pdf.showPage()
    pdf.save()

    deck = Presentation()
    for heading, lines in LESSONS:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = heading
        frame = slide.placeholders[1].text_frame
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(20)
    # Purely synthetic pixels: a small coordinate plot exercises embedded-image
    # previews without reading any private input or requiring an image service.
    diagram = Image.new('RGB', (480, 180), '#f4f7f6')
    draw = ImageDraw.Draw(diagram)
    draw.line([(30, 15), (30, 150), (450, 150)], fill='#344c49', width=3)
    curve = [(x, int(35 + 110 * ((x - 240) / 210) ** 2)) for x in range(40, 441)]
    draw.line(curve, fill='#236a86', width=4)
    for x, y in (curve[70], curve[125], curve[170], curve[200]):
        draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill='#d8673d')
    image_bytes = BytesIO()
    diagram.save(image_bytes, format='PNG')
    image_bytes.seek(0)
    deck.slides[0].shapes.add_picture(image_bytes, Inches(0.7), Inches(5.2), width=Inches(4.8))
    table = deck.slides[1].shapes.add_table(2, 2, Inches(0.6), Inches(5.4), Inches(8.7), Inches(1)).table
    for row, cells in enumerate([("Event 事件", "Probability 概率"), ("A given B", "0.4")]):
        for column, value in enumerate(cells):
            table.cell(row, column).text = value
    deck.save(output / "synthetic-foundations.pptx")

    document = Document()
    for heading, lines in LESSONS:
        document.add_heading(heading, level=1)
        for line in lines:
            document.add_paragraph(line)
    table = document.add_table(rows=2, cols=2)
    for row, cells in enumerate([("Matrix 矩阵", "Eigenvalues 特征值"), ("diag(2, 3)", "2, 3")]):
        for column, value in enumerate(cells):
            table.cell(row, column).text = value
    document.save(output / "synthetic-foundations.docx")

    for extension in ("txt", "md"):
        lines = []
        for heading, body in LESSONS:
            lines.append(("# " if extension == "md" else "") + heading)
            lines.extend(body)
            lines.extend(f"Practice {number}: {body[number % len(body)]}" for number in range(1, 25))
            lines.append("")
        (output / f"synthetic-foundations.{extension}").write_text("\n".join(lines), encoding="utf-8")


def main():
    project = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=project / "var/demo")
    args = parser.parse_args()
    output = args.output.resolve()
    if output == project or not output.is_relative_to(project):
        parser.error("Output must be a subdirectory within the project.")
    # Refuse destination-file symlinks as well as directory escapes.
    for kind in ("pdf", "pptx", "docx", "txt", "md"):
        target = output / f"synthetic-foundations.{kind}"
        if target.is_symlink():
            parser.error("Fixture files must stay within the project; remove the destination symlink.")
    generate(output)
    print(f"Generated 5 synthetic bilingual fixtures in {output}")


if __name__ == "__main__":
    main()
